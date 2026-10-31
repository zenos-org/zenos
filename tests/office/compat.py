#!/usr/bin/env python3
"""Office compatibility check: open each public file and save it again with an engine, then compare the saved file
with the original. Nothing here is a model or a guess: counts of what is in the files, text, cell values and
formulas, and pictures of the pages.

  compat.py CORPUS_DIR OUT_DIR [--engines zapps,libreoffice] [--limit N]

Engines:
  zapps        the ZENOS apps' own open/save code (apps/docs, apps/sheets, apps/slides), run by engines/zapps.bundle.mjs
  libreoffice  LibreOffice (soffice --headless --convert-to), the same file type in and out

Measures, original against saved:
  text         how much of the text survives (character pairs in common, 0..1)
  features     each kind of element counted in both files: tables, pictures, vertical text, ruby, distributed
               alignment, headers and footers (Word); merged cells, column widths, borders, print areas, page
               breaks, conditional formats, input rules, charts, formulas, macros (Excel); slides, pictures,
               charts, tables (PowerPoint)
  cells        Excel: values and formulas of every non-empty cell, the same
  pages        both files drawn by LibreOffice into page images: page counts and how much of the ink lands in the same place (0..1, ink in both / ink in either).
               The drawing is LibreOffice's, not Microsoft Office's: Microsoft Office itself is not used here.
Results: OUT_DIR/results.json and OUT_DIR/results.md.
"""
import json, os, re, shutil, subprocess, sys, tempfile, zipfile
from collections import Counter
from pathlib import Path

HERE = Path(__file__).resolve().parent
ZAPPS = HERE / "engines" / "zapps.bundle.mjs"

FEATURES = {
    "docx": {"tables": [r"<w:tbl>"], "pictures": [r"<w:drawing>", r"<v:imagedata"],
             "vertical text": [r"<w:textDirection w:val=\"(tbRl|tbRlV|btLr|tbLrV)\""], "ruby": [r"<w:ruby>"],
             "distributed": [r"<w:jc w:val=\"distribute\""],
             "headers/footers": ["@parts:word/(header|footer)\\d*\\.xml"]},
    "xlsx": {"merged cells": [r"<mergeCell "], "column widths": [r"<col "], "borders": ["@borders"],
             "print areas": [r"_xlnm\.Print_Area"], "page breaks": [r"<brk "],
             "conditional formats": [r"<conditionalFormatting"], "input rules": [r"<dataValidation "],
             "charts": ["@parts:xl/charts/chart\\d+\\.xml"], "formulas": [r"<f[ >]"],
             "macros": ["@parts:xl/vbaProject\\.bin"], "pictures": ["@parts:xl/media/"]},
    "pptx": {"slides": ["@parts:ppt/slides/slide\\d+\\.xml"], "pictures": [r"<p:pic>"],
             "charts": ["@parts:ppt/charts/chart\\d+\\.xml"], "tables": [r"<a:tbl>"]},
}
TEXT_PARTS = {"docx": r"word/(document|header\d*|footer\d*)\.xml", "xlsx": r"xl/(sharedStrings|worksheets/sheet\d+)\.xml",
              "pptx": r"ppt/slides/slide\d+\.xml"}
TEXT_TAGS = {"docx": r"<w:t[^>]*>([^<]*)</w:t>", "xlsx": r"<(?:t|v)[^>]*>([^<]*)</(?:t|v)>", "pptx": r"<a:t>([^<]*)</a:t>"}


LEGACY = ("doc", "xls", "ppt")          # the older binary formats: still most of what public offices publish


def kind(p):
    e = p.suffix.lower().lstrip(".")
    return {"docm": "docx", "xlsm": "xlsx", "pptm": "pptx"}.get(e, e)


def parts(path):
    try:
        with zipfile.ZipFile(path) as z:
            return {n: z.read(n) for n in z.namelist()}
    except (zipfile.BadZipFile, OSError):
        return None


def count(pp, k):
    out = {}
    xml = {n: v.decode("utf-8", "replace") for n, v in pp.items() if n.endswith(".xml") or n.endswith(".rels")}
    for name, pats in FEATURES[k].items():
        n = 0
        for pat in pats:
            if pat.startswith("@parts:"):
                n += sum(1 for p in pp if re.match(pat[7:], p))
            elif pat == "@borders":
                st = xml.get("xl/styles.xml", "")
                n += len(re.findall(r"<(left|right|top|bottom) style=", st))
            else:
                n += sum(len(re.findall(pat, v)) for v in xml.values())
        out[name] = n
    return out


def text(pp, k):
    s = []
    for n in sorted(pp):
        if re.fullmatch(TEXT_PARTS[k], n):
            s += re.findall(TEXT_TAGS[k], pp[n].decode("utf-8", "replace"))
    return re.sub(r"\s+", "", "".join(s))


def text_kept(a, b):
    if not a:
        return None
    ca, cb = Counter(a[i:i + 2] for i in range(len(a) - 1)), Counter(b[i:i + 2] for i in range(len(b) - 1))
    if not ca:
        return 1.0 if a in b else 0.0
    return round(sum((ca & cb).values()) / sum(ca.values()), 3)


def cells(orig, saved):
    try:
        import openpyxl
        wa, wb = openpyxl.load_workbook(orig), openpyxl.load_workbook(saved)
    except Exception:
        return None
    total = same = 0
    for i, sa in enumerate(wa.worksheets):
        sb = wb.worksheets[i] if i < len(wb.worksheets) else None
        for row in sa.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                total += 1
                v = sb[c.coordinate].value if sb is not None else None
                if str(v).lstrip("=") == str(c.value).lstrip("=") or (isinstance(v, (int, float)) and isinstance(c.value, (int, float)) and abs(v - c.value) < 1e-9):
                    same += 1
    return {"cells": total, "same": same, "kept": round(same / total, 3) if total else None}


PROFILE = None


def soffice(args, timeout=180):
    global PROFILE
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    prof = PROFILE = PROFILE or tempfile.mkdtemp(prefix="lo-")         # one profile: starting is much faster
    p = subprocess.Popen([exe, f"-env:UserInstallation=file://{prof}", "--headless", "--norestore"] + args,
                         stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, start_new_session=True)
    try:
        out, err = p.communicate(timeout=timeout)
    except subprocess.TimeoutExpired:
        os.killpg(p.pid, 9)            # soffice.bin too: a hung one would hold the profile and block every later file
        p.communicate()
        raise
    return subprocess.CompletedProcess(p.args, p.returncode, out, err)


def to_pdf(src, outdir):
    try:
        soffice(["--convert-to", "pdf", "--outdir", str(outdir), str(src)])
    except subprocess.TimeoutExpired:
        return None
    p = Path(outdir) / (Path(src).stem + ".pdf")
    return p if p.exists() else None


def pages(pdf, outdir, tag):
    subprocess.run(["pdftoppm", "-r", "30", "-gray", "-png", str(pdf), str(Path(outdir) / tag)], capture_output=True, timeout=300)
    return sorted(Path(outdir).glob(tag + "-*.png"))


def page_similarity(a, b):
    from PIL import Image
    import numpy as np
    n = max(len(a), len(b))
    if n == 0:
        return None
    s = 0.0
    for i in range(min(len(a), len(b))):
        ia, ib = Image.open(a[i]).convert("L"), Image.open(b[i]).convert("L")
        ib = ib.resize(ia.size)
        ka, kb = np.asarray(ia) < 200, np.asarray(ib) < 200          # the ink, not the white paper
        union = (ka | kb).sum()
        s += 1.0 if union == 0 else float((ka & kb).sum()) / float(union)
    return round(s / n, 3)


def run_engine(engine, src, out):
    if engine == "zapps":
        r = subprocess.run(["node", str(ZAPPS), str(src), str(out)], capture_output=True, text=True, timeout=180)
        return out.exists() and r.returncode == 0, r.stderr[-300:]
    if engine == "libreoffice":
        d = out.parent / ("lo-" + out.stem)
        d.mkdir(exist_ok=True)
        ext = src.suffix.lower().lstrip(".")
        flt = {"doc": "doc:MS Word 97", "xls": "xls:MS Excel 97", "ppt": "ppt:MS PowerPoint 97"}.get(ext, ext)
        r = soffice(["--convert-to", flt, "--outdir", str(d), str(src)])
        got = d / src.name
        if got.exists():
            shutil.move(str(got), out)
        return out.exists(), r.stderr[-300:]
    raise ValueError(engine)


def check(src, engine, work):
    k = kind(src)
    res = {"file": src.name, "kind": k, "engine": engine}
    out = work / f"{engine}-{src.name}"
    if engine == "zapps" and k in ("doc", "xls"):
        out = out.with_suffix(".docx" if k == "doc" else ".xlsx")      # the apps save in the newer format
    try:
        ok, err = run_engine(engine, src, out)
    except subprocess.TimeoutExpired:
        ok, err = False, "timeout"
    res["saved"] = ok
    if not ok:
        res["error"] = err
        return res
    if k in LEGACY:                   # no XML to count: compare the pages only
        pa = pb = None
    else:
        pa, pb = parts(src), parts(out)
    if k in LEGACY:
        pass
    elif pa is None or pb is None:
        res["error"] = "not a zip package"
        return res
    if pa is not None:
        fa, fb = count(pa, k), count(pb, k)
        res["features"] = {n: [fa[n], fb[n]] for n in fa}
        res["text"] = text_kept(text(pa, k), text(pb, k))
    if k == "xlsx":
        res["cells"] = cells(src, out)
    with tempfile.TemporaryDirectory() as t:
        t = Path(t)
        a, b = to_pdf(src, t / "a"), to_pdf(out, t / "b")
        if a and b:
            ia, ib = pages(a, t, "a"), pages(b, t, "b")
            res["pages"] = [len(ia), len(ib)]
            res["look"] = page_similarity(ia, ib)
    return res


def summary(results):
    lines = ["| engine | kind | files | saved | text kept | look alike | same page count |", "|---|---|---|---|---|---|---|"]
    feats = {}
    for (e, k) in sorted({(r["engine"], r["kind"]) for r in results}):
        rs = [r for r in results if r["engine"] == e and r["kind"] == k]
        sv = [r for r in rs if r.get("saved")]
        avg = lambda xs: round(sum(xs) / len(xs), 3) if xs else None
        tk = avg([r["text"] for r in sv if r.get("text") is not None])
        lk = avg([r["look"] for r in sv if r.get("look") is not None])
        pc = sum(1 for r in sv if r.get("pages") and r["pages"][0] == r["pages"][1])
        lines.append(f"| {e} | {k} | {len(rs)} | {len(sv)} | {tk} | {lk} | {pc}/{len(sv)} |")
        for r in sv:
            for n, (a, b) in (r.get("features") or {}).items():
                if a:
                    f = feats.setdefault((e, k, n), [0, 0])
                    f[0] += 1
                    f[1] += 1 if b >= a else 0
    lines += ["", "Files where an element was present, and how many kept all of it:", "",
              "| engine | kind | element | files with it | kept all |", "|---|---|---|---|---|"]
    for (e, k, n), (a, b) in sorted(feats.items()):
        lines.append(f"| {e} | {k} | {n} | {a} | {b} |")
    return "\n".join(lines) + "\n"


def main():
    a = sys.argv[1:]
    corpus, out = Path(a[0]), Path(a[1])
    engines = (a[a.index("--engines") + 1] if "--engines" in a else "zapps,libreoffice").split(",")
    limit = int(a[a.index("--limit") + 1]) if "--limit" in a else None
    out.mkdir(parents=True, exist_ok=True)
    files = sorted(p for p in corpus.iterdir() if kind(p) in FEATURES or kind(p) in LEGACY)[:limit]
    results = []
    with tempfile.TemporaryDirectory() as w:
        for f in files:
            for e in engines:
                r = check(f, e, Path(w))
                results.append(r)
                print(json.dumps(r, ensure_ascii=False), flush=True)
    (out / "results.json").write_text(json.dumps(results, ensure_ascii=False, indent=1))
    (out / "results.md").write_text(summary(results))
    print(summary(results))


if __name__ == "__main__":
    main()
